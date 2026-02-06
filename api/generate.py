from fastapi import APIRouter
import time, uuid
from docx import Document

from models.schemas import (
    DocumentGenerateRequest,
    DocumentGenerateResponse
)

router = APIRouter(prefix="/api/template", tags=["Templates"])


@router.post(
    "/generate",
    response_model=DocumentGenerateResponse
)
def generate_document(payload: DocumentGenerateRequest):
    start = time.time()
    trace_id = uuid.uuid4().hex[:8]

    template_id = payload.template_id
    data = payload.data

    template = template_json  # fetched from DB in real system
    doc = Document("templates/Hos_AFS.docx")

    replaced = 0

    # 1. Group table placeholders
    table_groups = group_placeholders_by_cell(template["placeholders"])

    # 2. Replace BODY placeholders
    for ph in template["placeholders"]:
        if ph["scope"] != "body":
            continue

        try:
            para = get_target_paragraph(doc, ph, trace_id)
            value = str(data.get(ph["key"], ""))

            if replace_in_paragraph(para, ph["syntax"], value, trace_id):
                replaced += 1
        except Exception:
            continue

    # 3. Replace TABLE placeholders
    for (t, r, c), ph_list in table_groups.items():
        cell = doc.tables[t].rows[r].cells[c]
        if replace_all_in_cell(cell, ph_list, data, trace_id):
            replaced += len(ph_list)

    output_path = f"outputs/{template_id}_{int(time.time())}.docx"
    doc.save(output_path)

    return {
        "status": "success",
        "replaced": replaced,
        "total": len(template["placeholders"]),
        "output": output_path,
        "generation_time_ms": int((time.time() - start) * 1000),
        "trace_id": trace_id
    }
