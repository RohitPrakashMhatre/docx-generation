from fastapi import APIRouter, UploadFile, File, Form, HTTPException
from docx import Document
import io 
from datetime import datetime 

from infrastructure.template_parser import (
    file_hash,
    parse_paragraphs,
    parse_tables
)
from models.upload import DocumentUploadResponse, DocumentUploadRequest


upl_router = APIRouter(prefix="/api/templates", tags=["Templates"])

@upl_router.post("/upload")
async def upload_templates(
        file: UploadFile=File(...),
        template_name: str=Form(...)
    ):

    file_bytes = await file.read()

    if not file_bytes:
        raise HTTPException(status_code=400, detail="Empty file uploaded")

    try:
        doc = Document(io.BytesIO(file_bytes))
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid DOCX file")

    placeholders = []
    placeholders.extend(parse_paragraphs(doc))
    placeholders.extend(parse_tables(doc))

    template_json = {
        "template_id": template_name,
        "meta": {
            "docx_hash": file_hash(file_bytes),
            "created_at": datetime.utcnow().isoformat(),
            "paragraph_count": len(doc.paragraphs),
            "table_count": len(doc.tables),
            "placeholder_count": len(placeholders)
        },
        "placeholders": placeholders
    }

    # db.templates.insert_one(template_json)
    

    return {
        "status": "parsed",
        "template_id": template_name,
        "placeholders_found": len(placeholders),
        "docx_hash": template_json["meta"]["docx_hash"],
        "patterns": template_json
    }

