from fastapi import APIRouter
import time, uuid, base64
from docx import Document
import datetime as datetime
import logging


from models.generate import (
    DocumentGenerateRequest,
    DocumentGenerateResponse
)
from domain.functions import evaluate_function_placeholder

gen_router = APIRouter(prefix="/api/template", tags=["Templates"])

@gen_router.post(
    "/generate-docx",
    response_model=DocumentGenerateResponse
)
def generate_document(payload: DocumentGenerateRequest):
    start = time.time()
    trace_id = uuid.uuid4().hex[:8]

    template_id = payload.template_id
    data = payload.data

    template_url = data.get("TemplateUrl")
    if not template_url:
        raise HTTPException(status_code=400, detail="TemplateUrl is required")

    # Download DOCX 
    try:
        response = requests.get(template_url, timeout=10)
        response.raise_for_status()
    except Exception as e:
        raise HTTPException(
            status_code=400,
            detail=f"Failed to download template: {str(e)}"
        )

    timestamp = f"{int(time.time() * 1000)}"
    original_name = template_url.split("/")[-1].replace("%20", "_")
    filename_docx = f"{timestamp}_edited_{original_name}"
    docx_file_path = f"outputs/{filename_docx}"

    with open(docx_file_path, "wb") as f:
        f.write(response.content)

    # Load DOCX 
    try:
        doc = Document(docx_file_path)
    except Exception:
        raise HTTPException(status_code=400, detail="Invalid DOCX file")

    template = template_json  # fetched from DB in real system
    replaced = 0

    # 1. Group table placeholders
    table_groups = group_placeholders_by_cell(template["placeholders"])

    # 2. Replace BODY placeholders
    for ph in template["placeholders"]:
        if ph["scope"] != "body":
            continue

        if ph["type"] == "function":
            value = evaluate_function_placeholder(ph["key"], data)
        elif ph["type"] == "clause_block":
            encode_text = data.get(ph["key"])
            decode_text = base64.b64decode(encode_text)
            value = decode_text.decode("utf-8")
        else:
            value = data.get(ph["key"])

        try:
            para = get_target_paragraph(doc, ph, trace_id)

            if replace_in_paragraph(para, ph["syntax"], str(value), trace_id):
                replaced += 1
        except Exception:
            continue

    # 3. Replace TABLE placeholders
    for (t, r, c), ph_list in table_groups.items():
        cell = doc.tables[t].rows[r].cells[c]
        if replace_all_in_cell(cell, ph_list, data, trace_id):
            replaced += len(ph_list)

    os.makedirs("outputs", exist_ok=True)
    output_path = f"outputs/{template_id}_{int(time.time())}.docx"
    doc.save(output_path)

    return {
        "status": "success",
        "replaced": replaced,
        "total": len(template["placeholders"]),
        "output": output_path,
        "generation_time_ms": int((time.time() - start) * 1000),
        "trace_id": trace_id
    }
